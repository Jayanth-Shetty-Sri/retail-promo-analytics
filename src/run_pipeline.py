"""
Rebuild the whole analysis in PostgreSQL and write this week's category report to Excel.

Run it from the repo folder, with the .venv active:

    python src/run_pipeline.py                 # rebuild every SQL step, then write the report
    python src/run_pipeline.py --report-only   # skip the rebuild, just write the report
    python src/run_pipeline.py --week 95       # report on a chosen week instead of the latest one
    python src/run_pipeline.py --load-raw      # reload the raw CSVs first (slow, asks for the password again)

The report lands in reports/weekly_category_report_week_<N>.xlsx.
"""
import argparse
import os
import subprocess
import sys
import time
from getpass import getpass
from pathlib import Path

import pandas as pd
import psycopg2
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import URL, create_engine

REPO = Path(__file__).resolve().parent.parent
SQL_STEPS = ["04_clean_views.sql", "05_uplift.sql", "06_cannibalisation.sql", "07_rfm.sql", "08_bi_views.sql"]
MIN_AVG_SALES = 100      # a category needs $100+ in an average week to count as a "mover"

BLUE, ORANGE, HEADER = "2A78D6", "EB6834", "14213D"


# ---------------------------------------------------------------- 1. rebuild

def rebuild(conn):
    """Run the SQL scripts in order, keeping the campaign results from notebook 02."""
    cur = conn.cursor()

    # 05_uplift.sql drops the whole analysis schema, which also deletes the campaign
    # results notebook 02 saved there. I park them in a side table and put them back,
    # so I don't have to rerun the slow difference-in-differences notebook every week.
    cur.execute("SELECT to_regclass('analysis.campaign_did') IS NOT NULL")
    has_campaigns = cur.fetchone()[0]
    if has_campaigns:
        cur.execute("DROP TABLE IF EXISTS public.campaign_did_keep;"
                    "CREATE TABLE public.campaign_did_keep AS SELECT * FROM analysis.campaign_did;")
        conn.commit()

    for name in SQL_STEPS:
        if name == "08_bi_views.sql":
            if not has_campaigns:
                sys.exit("Stopped before 08_bi_views.sql: analysis.campaign_did doesn't exist yet.\n"
                         "Run notebooks/02_campaign_impact.ipynb once, then run this script again.")
            cur.execute("CREATE TABLE analysis.campaign_did AS SELECT * FROM public.campaign_did_keep;"
                        "DROP TABLE public.campaign_did_keep;")
            conn.commit()

        start = time.time()
        print(f"  running {name} ...", end=" ", flush=True)
        cur.execute((REPO / "sql" / name).read_text(encoding="utf-8"))
        conn.commit()
        print(f"done in {time.time() - start:.0f}s")


# ---------------------------------------------------------------- 2. the numbers

def pct(new, old):
    """Percentage change, or None when there's nothing to compare with."""
    return (new / old - 1) if old else None


def weekly_numbers(engine, week):
    """Category and department sales for the report week, with the comparisons."""
    cat = pd.read_sql(f"""
        SELECT department, commodity_desc AS category, week_no, sales, discount
        FROM bi.fact_category_week
        WHERE week_no BETWEEN {week - 52} AND {week}""", engine)

    sales = cat.pivot_table(index=["department", "category"], columns="week_no",
                            values="sales", aggfunc="sum", fill_value=0)
    disc = cat[cat["week_no"] == week].groupby(["department", "category"])["discount"].sum()

    def col(w):
        return sales[w] if w in sales.columns else 0

    out = pd.DataFrame(index=sales.index)
    out["sales_this_week"] = col(week)
    out["sales_last_week"] = col(week - 1)
    out["avg_last_4_weeks"] = sum(col(w) for w in range(week - 4, week)) / 4
    out["same_week_last_year"] = col(week - 52)
    out["vs_last_week"] = [pct(a, b) for a, b in zip(out["sales_this_week"], out["sales_last_week"])]
    out["vs_4_week_avg"] = [pct(a, b) for a, b in zip(out["sales_this_week"], out["avg_last_4_weeks"])]
    out["vs_last_year"] = [pct(a, b) for a, b in zip(out["sales_this_week"], out["same_week_last_year"])]
    out["discount_pct"] = (disc.reindex(out.index).fillna(0) / out["sales_this_week"]).where(out["sales_this_week"] > 0)
    out["share_of_sales"] = out["sales_this_week"] / out["sales_this_week"].sum()
    out = out.reset_index()

    money = ["sales_this_week", "sales_last_week", "avg_last_4_weeks", "same_week_last_year"]
    dept = out.groupby("department")[money].sum()
    dept["vs_last_week"] = [pct(a, b) for a, b in zip(dept["sales_this_week"], dept["sales_last_week"])]
    dept["vs_4_week_avg"] = [pct(a, b) for a, b in zip(dept["sales_this_week"], dept["avg_last_4_weeks"])]
    dept["vs_last_year"] = [pct(a, b) for a, b in zip(dept["sales_this_week"], dept["same_week_last_year"])]
    dept["share_of_sales"] = dept["sales_this_week"] / dept["sales_this_week"].sum()
    dept = dept.reset_index().sort_values("sales_this_week", ascending=False)

    return out.sort_values("sales_this_week", ascending=False), dept


def headline(engine, week):
    """The four executive numbers for the report week and the week before."""
    kpi = pd.read_sql(f"""
        SELECT week_no,
               SUM(basket_value)              AS sales,
               COUNT(DISTINCT household_key)  AS households,
               COUNT(*)                       AS trips,
               SUM(basket_value) / COUNT(*)   AS basket_value
        FROM bi.fact_baskets
        WHERE week_no IN ({week - 1}, {week})
        GROUP BY week_no""", engine).set_index("week_no")
    rows = []
    for metric in ["sales", "households", "trips", "basket_value"]:
        this = kpi[metric].get(week, 0)
        last = kpi[metric].get(week - 1, 0)
        rows.append([metric.replace("_", " ").capitalize(), this, last, pct(this, last)])
    return pd.DataFrame(rows, columns=["metric", "this_week", "last_week", "change"])


# ---------------------------------------------------------------- 3. the Excel file

def style_sheet(ws, money_cols, pct_cols, change_cols, first_row=1):
    """Header colours, number formats, column widths, and blue/orange for ups and downs."""
    header = [c.value for c in ws[first_row]]
    for cell in ws[first_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, name in enumerate(header, start=1):
        letter = ws.cell(row=first_row, column=i).column_letter
        ws.column_dimensions[letter].width = max(14, min(32, len(str(name)) + 4))
        cells = f"{letter}{first_row + 1}:{letter}{ws.max_row}"
        fmt = "$#,##0" if name in money_cols else "0.0%" if name in pct_cols else None
        if fmt:
            for row in ws[cells]:
                row[0].number_format = fmt
        if name in change_cols:
            ws.conditional_formatting.add(cells, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=BLUE)))
            ws.conditional_formatting.add(cells, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=ORANGE)))
    ws.freeze_panes = ws.cell(row=first_row + 1, column=1)


def write_report(path, week, kpi, dept, cat):
    money = ["sales_this_week", "sales_last_week", "avg_last_4_weeks", "same_week_last_year"]
    change = ["vs_last_week", "vs_4_week_avg", "vs_last_year"]
    pcts = change + ["discount_pct", "share_of_sales"]

    steady = cat[cat["avg_last_4_weeks"] >= MIN_AVG_SALES].copy()
    steady["change_vs_4_week_avg"] = steady["sales_this_week"] - steady["avg_last_4_weeks"]
    rising = steady[steady["change_vs_4_week_avg"] > 0].nlargest(10, "change_vs_4_week_avg")
    falling = steady[steady["change_vs_4_week_avg"] < 0].nsmallest(10, "change_vs_4_week_avg")
    mover_cols = ["department", "category", "sales_this_week", "avg_last_4_weeks", "change_vs_4_week_avg", "vs_4_week_avg"]

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        kpi.to_excel(xl, sheet_name="Summary", index=False, startrow=2)
        dept.to_excel(xl, sheet_name="Departments", index=False)
        cat.to_excel(xl, sheet_name="Categories", index=False)
        rising[mover_cols].to_excel(xl, sheet_name="Rising", index=False)
        falling[mover_cols].to_excel(xl, sheet_name="Falling", index=False)

        ws = xl.sheets["Summary"]
        ws["A1"] = f"Brightbasket weekly category report: week {week}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "Loyalty-household sales, merchandise only. Compared with the week before."
        style_sheet(ws, [], ["change"], ["change"], first_row=3)
        for r in range(4, 8):
            ws.cell(row=r, column=2).number_format = ws.cell(row=r, column=3).number_format = (
                "$#,##0.00" if ws.cell(row=r, column=1).value in ("Sales", "Basket value") else "#,##0")
        ws.column_dimensions["A"].width = 18

        style_sheet(xl.sheets["Departments"], money, pcts, change)
        style_sheet(xl.sheets["Categories"], money, pcts, change)
        for name in ["Rising", "Falling"]:
            style_sheet(xl.sheets[name], money + ["change_vs_4_week_avg"], ["vs_4_week_avg"],
                        ["change_vs_4_week_avg", "vs_4_week_avg"])


# ---------------------------------------------------------------- main

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--report-only", action="store_true", help="skip the SQL rebuild")
    parser.add_argument("--load-raw", action="store_true", help="reload the raw CSVs first")
    parser.add_argument("--week", type=int, help="week number to report on (default: latest complete week)")
    args = parser.parse_args()

    os.chdir(REPO)
    started = time.time()

    if args.load_raw:
        print("1. Reloading the raw CSVs")
        subprocess.run([sys.executable, "src/load_raw.py"], check=True)

    password = os.environ.get("PG_PASSWORD") or getpass("PostgreSQL password: ")
    db = dict(host="localhost", port=5432, dbname="retail_promo", user="postgres", password=password)

    if not args.report_only:
        print("2. Rebuilding the clean layer, the analysis and the dashboard views")
        with psycopg2.connect(**db) as conn:
            rebuild(conn)

    engine = create_engine(URL.create("postgresql+psycopg2", username=db["user"], password=password,
                                      host=db["host"], port=db["port"], database=db["dbname"]))
    week = args.week or int(pd.read_sql(
        "SELECT MAX(week_no) AS w FROM bi.dim_week WHERE complete_week", engine)["w"].iloc[0])

    print(f"3. Writing the report for week {week}")
    cat, dept = weekly_numbers(engine, week)
    kpi = headline(engine, week)
    Path("reports").mkdir(exist_ok=True)
    path = Path("reports") / f"weekly_category_report_week_{week}.xlsx"
    write_report(path, week, kpi, dept, cat)

    sales = kpi.set_index("metric").loc["Sales"]
    change = f"{sales['change']:+.1%}" if sales["change"] is not None else "n/a"
    print(f"\nWeek {week}: sales ${sales['this_week']:,.0f} ({change} vs last week)")
    print(f"Saved {path} in {time.time() - started:.0f}s")


if __name__ == "__main__":
    main()
